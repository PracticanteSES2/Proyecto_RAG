"""
Ingesta de la documentación viva (.md) y reindexación por grupo.

Ejecutar:  python -m unittest tests.docs_md.test_ingesta

Cada prueba arma un proyecto temporal (data/raw, data/catalog, Qdrant local)
con una COPIA del fixture tests/docs_md/fixtures/tablero_tiempos_urgencias
(manifest + md/*.md, las páginas que haya), un .docx falso en ese mismo grupo
(debe ignorarse: si se abriera, python-docx fallaría) y otro grupo
«tablero_otro» con un .docx real (el flujo .docx sigue igual). Nunca toca el
data/ real ni la carpeta del fixture. DOCS_MD_FIXTURE=<carpeta> usa otro
fixture con el mismo formato. TestRaizMarkdown prueba los .md fuera de
data/raw (md_root / DOCS_MD_ROOT, manifest junto a los .md o en data/raw).

Embeddings:
- Reindexación (aislamiento de grupos, errores): embedder FALSO inyectable
  (bolsa de palabras con hash, 64 dimensiones). Es instantáneo y lo que se
  prueba es qué puntos se reemplazan, no la calidad semántica.
- Recuperación (semáforo de triage II antes y después de editar el .md): el
  modelo REAL de embedding_indexer (paraphrase-multilingual-MiniLM-L12-v2) y
  el HybridRetriever del proyecto; con el modelo en la caché local tarda
  ~10-60 s (sobre todo importar torch). Si el modelo no está en la caché
  local la prueba se omite (HF_HUB_OFFLINE=1: nunca descarga).
"""

import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.dont_write_bytecode = True
os.environ.setdefault("HF_HUB_OFFLINE", "1")

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.docs_md import formato, repositorio  # noqa: E402
from src.docs_md.indexacion import reindexar_grupo, rutas_datos  # noqa: E402

FIXTURE = Path(os.getenv("DOCS_MD_FIXTURE") or Path(__file__).parent / "fixtures" / "tablero_tiempos_urgencias")
GRUPO = "tablero_tiempos_urgencias"
OTRO = "tablero_otro"
MODELO = "TABLERO TIEMPOS URGENCIAS"
DOCX_VIEJO = "3 Documentacion Tablero Tiempos Urgencias V2.docx"

# Páginas que la prueba agrega en la COPIA si el fixture aún no las tiene.
PAGINA_MEDIDAS = """---
tablero: TABLERO TIEMPOS URGENCIAS
pagina: PRUEBA MEDIDAS
nombre_funcional: Página de prueba con medidas citadas
bloque: Histórico
actualizado: 2026-10-20
version: 1
---

# PRUEBA MEDIDAS

## Propósito
Página de prueba para el enlace de medidas documentadas.

## Indicadores
- **Total de triages**: Cantidad de triages del periodo. Medida [TOTAL_TRIAGES].
- **Indicador sin medida en el modelo**: Medida [MEDIDA_QUE_NO_EXISTE] sobre CONSULTAS_DIA[TIEMPO_ESPERA_CONS].

## Historial de cambios
- **2026-10-20 · v1**: Documentación inicial.
"""

PAGINAS_EXTRA = {
    "_general.md": """---
tablero: TABLERO TIEMPOS URGENCIAS
pagina: GENERAL
nombre_funcional: Tiempos de urgencias: visión general
bloque: General
actualizado: 2026-10-20
version: 1
---

# GENERAL

## Propósito
Mide la oportunidad de la atención en urgencias en cada etapa del recorrido del paciente.

## Definiciones
- **Oportunidad**: Minutos entre dos eventos de la atención del paciente.

## Historial de cambios
- **2026-10-20 · v1**: Documentación inicial.
""",
    "interconsultas_dia.md": """---
tablero: TABLERO TIEMPOS URGENCIAS
pagina: INTERCONSULTAS DIA
bloque: Día
actualizado: 2026-10-20
version: 1
---

# INTERCONSULTAS DIA

## Propósito
Muestra las interconsultas del día pendientes y respondidas, con semáforo.

## Reglas de negocio
- **Semáforo de interconsultas**: Verde de 0 a 12 horas, amarillo de más de 12 hasta 24 y rojo de más de 24.

## Historial de cambios
- **2026-10-20 · v1**: Documentación inicial.
""",
    "interconsultas_historico.md": """---
tablero: TABLERO TIEMPOS URGENCIAS
pagina: INTERCONSULTAS HISTORICO
bloque: Histórico
actualizado: 2026-10-20
version: 1
---

# INTERCONSULTAS HISTORICO

## Propósito
Muestra el tiempo promedio de respuesta de las interconsultas por año y mes.

## Indicadores
- **Tiempo de respuesta de interconsulta**: Horas entre la solicitud y la respuesta. Columna INTERCONSULTA[TIEMPO_RESPUESTA].

## Historial de cambios
- **2026-10-20 · v1**: Documentación inicial.
""",
}

CATALOGO_TECNICO = {
    "workspace": "Gestion Clinica",
    "semantic_model": MODELO,
    "tables": [
        {"name": "CONSULTAS_DIA", "storage_mode": "DirectQuery", "measures": []},
        {"name": "OPORTUNIDAD_TRIAGE", "storage_mode": "Import",
         "measures": [{"name": "TOTAL_TRIAGES", "expression": "COUNTROWS(OPORTUNIDAD_TRIAGE)"}]},
    ],
    "relationships": [],
}


# ============================================================
# Dobles y utilidades
# ============================================================

class EmbedderFalso:
    """Bolsa de palabras con hash: determinista e instantáneo."""

    DIM = 64

    def get_sentence_embedding_dimension(self):
        return self.DIM

    def encode(self, textos, **_):
        unico = isinstance(textos, str)
        vectores = []
        for texto in [textos] if unico else textos:
            vector = [0.0] * self.DIM
            for palabra in re.findall(r"\w+", str(texto).lower()):
                vector[int(hashlib.md5(palabra.encode("utf-8")).hexdigest(), 16) % self.DIM] += 1.0
            norma = math.sqrt(sum(v * v for v in vector)) or 1.0
            vectores.append([v / norma for v in vector])
        return vectores[0] if unico else vectores


def _paginas_fixture():
    return sorted((FIXTURE / "md").glob("*.md"))


def _crear_docx(ruta: Path):
    from docx import Document

    documento = Document()
    documento.add_paragraph("DOCUMENTACIÓN TABLERO OTRO")
    documento.add_paragraph("Este tablero muestra la ocupación de camas por servicio.")
    documento.add_paragraph("DESCRIPCIÓN DE LOS VISUALES")
    documento.add_paragraph("Tabla de ocupación")
    documento.add_paragraph("Muestra las camas ocupadas por servicio y piso.")
    documento.save(ruta)


def crear_proyecto(raiz: Path) -> Path:
    """Proyecto temporal: copia del fixture + páginas de prueba + otro grupo."""
    raw = raiz / "data" / "raw"
    grupo = raw / GRUPO
    shutil.copytree(FIXTURE, grupo)
    (grupo / DOCX_VIEJO).write_bytes(b"no es un docx: si se abre, falla")
    md = grupo / "md"
    (md / "prueba_medidas.md").write_text(PAGINA_MEDIDAS, encoding="utf-8")
    for nombre, texto in PAGINAS_EXTRA.items():
        if not (md / nombre).exists():
            (md / nombre).write_text(texto, encoding="utf-8")

    otro = raw / OTRO
    otro.mkdir(parents=True)
    (otro / "manifest.json").write_text(json.dumps({
        "workspace": "Gestion Clinica", "report": "TABLERO OTRO", "semantic_model": "TABLERO OTRO",
        "semantic_model_key": "tablero_otro", "source_group": OTRO, "default_dashboard": "OTRO",
    }), encoding="utf-8")
    _crear_docx(otro / "1 Documentacion Otro.docx")

    catalogo = raiz / "data" / "catalog"
    catalogo.mkdir(parents=True)
    (catalogo / "tablero_tiempos_urgencias_rag.json").write_text(
        json.dumps(CATALOGO_TECNICO, ensure_ascii=False), encoding="utf-8")
    return raw


def _paginas_esperadas(raw: Path) -> dict:
    """{nombre del tablero normalizado: Documento} de las páginas de la copia."""
    esperadas = {}
    for archivo in sorted((raw / GRUPO / "md").glob("*.md")):
        doc = formato.parsear(archivo.read_text(encoding="utf-8"))
        nombre = MODELO if doc.meta.get("bloque") == "General" else doc.pagina
        esperadas[nombre] = doc
    return esperadas


def _pipeline_sin_indice(raiz: Path):
    """Extracción + normalización + catálogo + chunks (como ingest_rag, sin Qdrant)."""
    from src.ingestion.document_parser import process_all_documents
    from src.rag.chunk_builder import build_chunks
    from src.semantic.document_normalizer import normalize_all_documents
    from src.semantic.semantic_merge import build_semantic_catalog

    rutas = rutas_datos(raiz)
    extraccion = process_all_documents(rutas["raw"], rutas["processed"])
    normalizacion = normalize_all_documents(rutas["processed"], rutas["semantic_documents"])
    catalogo = build_semantic_catalog(None, rutas["semantic_documents"], rutas["catalogo"], project_root=raiz)
    return extraccion, normalizacion, catalogo, build_chunks(catalogo)


def _silencio():
    import contextlib
    import io

    return contextlib.redirect_stdout(io.StringIO())


def _puntos(cliente, coleccion):
    puntos, offset = [], None
    while True:
        lote, offset = cliente.scroll(coleccion, limit=256, offset=offset, with_payload=True, with_vectors=False)
        puntos.extend(lote)
        if offset is None:
            return puntos


class _ConProyecto(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="docs_md_ingesta_"))
        self.raw = crear_proyecto(self.tmp)
        self.rutas = rutas_datos(self.tmp)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


# ============================================================
# 1-2. Ingesta .md -> JSON normalizado -> catálogo -> chunks
# ============================================================

class TestIngestaMarkdown(_ConProyecto):

    def test_md_reemplaza_docx_y_produce_un_tablero_por_pagina(self):
        with _silencio():
            extraccion, normalizacion, catalogo, _ = _pipeline_sin_indice(self.tmp)
        self.assertEqual(extraccion["errors"], [])
        self.assertEqual(normalizacion["errors"], [])
        self.assertIn(GRUPO, extraccion["markdown_groups"])
        self.assertEqual([Path(p).name for p in extraccion["ignored_docx"]], [DOCX_VIEJO])

        esperadas = _paginas_esperadas(self.raw)
        # Solo JSON del .md en el grupo; el .docx del otro grupo sí se procesa.
        semanticos = sorted((self.rutas["semantic_documents"] / GRUPO).rglob("*.json"))
        self.assertEqual(len(semanticos), len(esperadas))
        self.assertTrue(all(p.parent.name == "md" for p in semanticos))
        self.assertTrue((self.rutas["semantic_documents"] / OTRO / "1 Documentacion Otro.json").exists())

        tableros = {t["name"]: t for t in catalogo["dashboards"] if t["source_group"] == GRUPO}
        self.assertEqual(set(tableros), set(esperadas))

        consultas = tableros["CONSULTAS DIA"]
        doc = esperadas["CONSULTAS DIA"]
        self.assertEqual(consultas["aliases"], [])
        self.assertEqual(consultas["source_file"], "consultas_dia.md")
        self.assertEqual(consultas["semantic_model"], MODELO)
        self.assertEqual(consultas["workspace"], "Gestion Clinica")
        self.assertTrue(consultas["manifest_path"].endswith("manifest.json"))
        self.assertIn(doc.proposito.split(".")[0], consultas["description"])
        self.assertIn(doc.meta["nombre_funcional"], consultas["description"])
        # «CONSULTAS REALIZADAS» es un título de visual, no SQL.
        self.assertEqual(consultas["sql_queries"], [])
        for seccion, campo in (("Visuales", "visuals"), ("Indicadores", "indicators"),
                               ("Reglas de negocio", "business_rules"), ("Fuentes de datos", "data_sources"),
                               ("Observaciones", "notes")):
            self.assertEqual(consultas[campo], [f"{e.nombre}: {e.texto}" for e in doc.entradas(seccion)], campo)
        self.assertIn("Semáforo triage II: Verde de 0 a 15 minutos, amarillo de 16 a 30 y rojo de 31 o más.",
                      consultas["business_rules"])
        historial = [s for s in consultas["other_sections"] if s["heading"] == "Historial de cambios"]
        self.assertEqual(len(historial), len(doc.entradas("Historial de cambios")))
        self.assertTrue(historial[0]["text"].startswith("2026-10-20 · v1: "))
        # La tabla citada en el .md se relaciona con el modelo técnico.
        self.assertIn("CONSULTAS_DIA", [t["name"] for t in consultas["powerbi_tables"]])

        # _general: se llama como el tablero (no «GENERAL») y lleva los alias del manifest.
        general = tableros[MODELO]
        self.assertIn("TIEMPOS DE URGENCIAS", general["aliases"])
        self.assertIn("TABLERO TIEMPOS DE URGENCIAS", general["aliases"])
        self.assertNotIn("GENERAL", [a.upper() for a in general["aliases"]])
        self.assertTrue(general["definitions"])

        self.assertIn("INTERCONSULTAS DIA", tableros)
        self.assertIn("INTERCONSULTAS HISTORICO", tableros)

    def test_medidas_documentadas_se_asocian_al_modelo(self):
        with _silencio():
            _, _, catalogo, chunks = _pipeline_sin_indice(self.tmp)
        prueba = next(t for t in catalogo["dashboards"] if t["name"] == "PRUEBA MEDIDAS")
        self.assertEqual([m["documented"]["name"] for m in prueba["matched_measures"]], ["TOTAL_TRIAGES"])
        self.assertEqual(prueba["matched_measures"][0]["match_type"], "exact")
        self.assertEqual([m["name"] for m in prueba["unmatched_documented_measures"]], ["MEDIDA_QUE_NO_EXISTE"])
        # CONSULTAS_DIA[TIEMPO_ESPERA_CONS] es una columna, no una medida.
        nombres = [m["name"] for m in prueba["documented_measures"]]
        self.assertNotIn("TIEMPO_ESPERA_CONS", nombres)

        medida = [c for c in chunks if c["chunk_type"] == "measure" and c["metadata"].get("dashboard") == "PRUEBA MEDIDAS"]
        self.assertEqual(len(medida), 1)
        self.assertIn("COUNTROWS(OPORTUNIDAD_TRIAGE)", medida[0]["text"])
        documentada = [c for c in chunks if c["chunk_type"] == "documented_measure"
                       and c["metadata"].get("dashboard") == "PRUEBA MEDIDAS"]
        self.assertEqual([c["metadata"]["measure"] for c in documentada], ["MEDIDA_QUE_NO_EXISTE"])

    def test_un_chunk_por_entrada(self):
        with _silencio():
            _, _, _, chunks = _pipeline_sin_indice(self.tmp)
        esperadas = _paginas_esperadas(self.raw)
        del_grupo = [c for c in chunks if c["metadata"]["source_group"] == GRUPO]
        self.assertFalse([c for c in del_grupo if c["chunk_type"] == "sql_query"])
        self.assertEqual(len({c["id"] for c in chunks}), len(chunks))

        tipos = {"Visuales": "visual", "Indicadores": "indicator", "Reglas de negocio": "business_rule",
                 "Fuentes de datos": "data_source", "Observaciones": "note", "Definiciones": "definition"}
        for nombre, doc in esperadas.items():
            de_pagina = [c for c in del_grupo if c["metadata"]["dashboard"] == nombre]
            self.assertEqual(sum(c["chunk_type"] == "dashboard_overview" for c in de_pagina), 1, nombre)
            for seccion, tipo in tipos.items():
                self.assertEqual(sum(c["chunk_type"] == tipo for c in de_pagina), len(doc.entradas(seccion)),
                                 f"{nombre} / {seccion}")
            historial = [c for c in de_pagina if c["chunk_type"] == "document_section"
                         and c["metadata"].get("section") == "Historial de cambios"]
            self.assertEqual(len(historial), len(doc.entradas("Historial de cambios")), nombre)

        regla = [c for c in del_grupo if c["metadata"]["dashboard"] == "CONSULTAS DIA"
                 and c["chunk_type"] == "business_rule" and "Semáforo triage II:" in c["text"]]
        self.assertEqual(len(regla), 1)
        self.assertEqual(regla[0]["text"], "Tablero: CONSULTAS DIA Regla de negocio: Semáforo triage II: "
                                           "Verde de 0 a 15 minutos, amarillo de 16 a 30 y rojo de 31 o más.")
        self.assertEqual(regla[0]["metadata"]["source_file"], "consultas_dia.md")
        # El flujo .docx sigue igual en los demás grupos.
        otro = [c for c in chunks if c["metadata"]["source_group"] == OTRO]
        self.assertIn("dashboard_overview", {c["chunk_type"] for c in otro})
        self.assertIn("visual", {c["chunk_type"] for c in otro})

    def test_sin_limpieza_previa_no_quedan_json_del_docx(self):
        """--no-clean-derived: el JSON de una corrida anterior del .docx se elimina."""
        viejo = self.rutas["semantic_documents"] / GRUPO / "3 Documentacion Tablero Tiempos Urgencias V2.json"
        viejo_procesado = self.rutas["processed"] / GRUPO / "3 Documentacion Tablero Tiempos Urgencias V2.json"
        for ruta in (viejo, viejo_procesado):
            ruta.parent.mkdir(parents=True, exist_ok=True)
            ruta.write_text(json.dumps({"semantic_model": MODELO, "dashboards": [{"name": "VIEJO"}]}),
                            encoding="utf-8")
        with _silencio():
            _, _, catalogo, _ = _pipeline_sin_indice(self.tmp)
        self.assertFalse(viejo.exists())
        self.assertFalse(viejo_procesado.exists())
        self.assertNotIn("VIEJO", [t["name"] for t in catalogo["dashboards"]])

    def test_query_plan_builder_y_enrutador_reconocen_las_paginas(self):
        from src.semantic.query_plan_builder import QueryPlanBuilder, normalize_text
        from src.semantic.source_model_router import SourceModelRouter

        with _silencio():
            _, _, catalogo, _ = _pipeline_sin_indice(self.tmp)
        rag = self.tmp / "data" / "rag"
        rag.mkdir(parents=True, exist_ok=True)
        espera = {
            "metric_id": "u_espera_consulta", "label": "Tiempo de espera consulta", "measure": "PROM_ESPERA",
            "aliases": [], "semantic_model": MODELO, "report": MODELO, "reports": [MODELO],
            "table": "CONSULTAS_DIA", "dax_expression": "AVERAGE(CONSULTAS_DIA[TIEMPO_ESPERA_CONS])",
            "validation_status": "approved",
            "appearances": [{"report": MODELO, "page_display_name": "CONSULTAS DIA", "visual_id": "v1",
                             "visual_type": "card", "visual_title": "Tiempo de espera consulta"}],
        }
        triages = {
            "metric_id": "u_total_triages", "label": "TOTAL_TRIAGES", "measure": "TOTAL_TRIAGES", "aliases": [],
            "semantic_model": MODELO, "report": None, "reports": [], "table": "OPORTUNIDAD_TRIAGE",
            "dax_expression": "[TOTAL_TRIAGES]", "validation_status": "approved", "appearances": [],
        }
        (rag / "master_metrics.json").write_text(json.dumps({"metrics": [espera, triages]}), encoding="utf-8")
        (rag / "visual_metrics_catalog.json").write_text(json.dumps({"schema_version": 2, "reports": [{
            "report": MODELO, "semantic_model": MODELO, "source_group": GRUPO,
            "pages": [{"page_name": "p1", "page_display_name": "CONSULTAS DIA", "filters": {}, "visuals": [
                {"visual_id": "v1", "visual_type": "card", "fields": [
                    {"role": "Values", "kind": "measure", "table": "CONSULTAS_DIA", "measure": "PROM_ESPERA",
                     "query_ref": "CONSULTAS_DIA.PROM_ESPERA"}]}]}],
        }]}), encoding="utf-8")
        registro = self.tmp / "data" / "catalog" / "source_registry.json"
        registro.write_text(json.dumps({"sources": [{
            "source_group": GRUPO, "report": MODELO, "semantic_model": MODELO, "default_dashboard": MODELO,
            "aliases": ["TIEMPOS DE URGENCIAS"], "technical_catalog": "data/catalog/tablero_tiempos_urgencias_rag.json",
        }], "models": []}), encoding="utf-8")

        router = SourceModelRouter(registro, rag / "visual_metrics_catalog.json", project_root=self.tmp)
        router.register_documentation_dashboards(
            [(t["name"], t["source_group"]) for t in catalogo["dashboards"]])
        builder = QueryPlanBuilder(rag / "master_metrics.json", rag / "visual_metrics_catalog.json",
                                   router, None, project_root=self.tmp)

        documentados = {t["name"] for t in builder._semantic_documents_index()[normalize_text(MODELO)]}
        self.assertTrue({"CONSULTAS DIA", "INTERCONSULTAS DIA", "INTERCONSULTAS HISTORICO", MODELO} <= documentados)

        # La documentación de la página se asocia al indicador que está en ella.
        textos = " ".join(builder._metric_documentation(espera))
        self.assertIn("Semáforo triage II: Verde de 0 a 15 minutos", textos)
        # La medida citada como [TOTAL_TRIAGES] liga el indicador a su página documentada.
        paginas = builder._page_index()
        miembros = paginas[(normalize_text(MODELO), normalize_text("PRUEBA MEDIDAS"))]["members"]
        self.assertIn("u_total_triages", [m["metric"]["metric_id"] for m in miembros.values()])
        self.assertTrue(paginas[(normalize_text(MODELO), normalize_text("CONSULTAS DIA"))]["visual"])

        # Enrutador: páginas con «DIA»/«HISTORICO» y el tablero completo.
        for pregunta, pagina in (("tiempos en la página de interconsultas dia", "INTERCONSULTAS DIA"),
                                 ("tiempos en la página de interconsultas historico", "INTERCONSULTAS HISTORICO")):
            menciones = router.board_mentions(pregunta)
            self.assertTrue(menciones and menciones[0]["known"], pregunta)
            self.assertIn(pagina, [c["dashboard"] for c in menciones[0]["contexts"]], pregunta)
        menciones = router.board_mentions("cuántos triages hay en el tablero de tiempos de urgencias")
        self.assertTrue(menciones[0]["known"])
        self.assertEqual({c["semantic_model"] for c in menciones[0]["contexts"]}, {MODELO})
        # _general no se registra como una página «GENERAL».
        self.assertFalse(any(m["known"] for m in router.board_mentions("datos del tablero general")))
        self.assertNotIn("general", [normalize_text(p["name"]) for p in router.known_pages()])


# ============================================================
# 3-4. reindexar_grupo (embedder falso, Qdrant temporal)
# ============================================================

class TestReindexarGrupo(_ConProyecto):

    def _cliente(self):
        from qdrant_client import QdrantClient

        self.rutas["qdrant"].mkdir(parents=True, exist_ok=True)
        cliente = QdrantClient(path=str(self.rutas["qdrant"]))
        self.addCleanup(cliente.close)
        return cliente

    def _por_grupo(self, cliente):
        from src.rag.embedding_indexer import COLLECTION_NAME

        grupos = {}
        for punto in _puntos(cliente, COLLECTION_NAME):
            grupos.setdefault(punto.payload["source_group"], {})[punto.id] = punto.payload["text"]
        return grupos

    def test_reemplaza_solo_el_grupo(self):
        cliente = self._cliente()
        embedder = EmbedderFalso()
        otro = reindexar_grupo(OTRO, cliente, project_root=self.tmp, embedder=embedder, verbose=False)
        self.assertEqual(otro["errores"], [])
        primero = reindexar_grupo(GRUPO, cliente, project_root=self.tmp, embedder=embedder, verbose=False)

        for clave in ("source_group", "documentos", "chunks", "segundos", "errores"):
            self.assertIn(clave, primero)
        self.assertEqual(primero["errores"], [])
        self.assertEqual(primero["source_group"], GRUPO)
        self.assertEqual(primero["documentos"], len(list((self.raw / GRUPO / "md").glob("*.md"))))
        self.assertEqual(primero["tableros"], primero["documentos"])
        self.assertEqual([Path(p).name for p in primero["docx_ignorados"]], [DOCX_VIEJO])
        self.assertGreater(primero["segundos"], 0)

        antes = self._por_grupo(cliente)
        self.assertEqual(len(antes[GRUPO]), primero["chunks"])
        self.assertEqual(len(antes[OTRO]), otro["chunks"])

        # Cambio aprobado: se modifica y se elimina una entrada de CONSULTAS DIA.
        doc = repositorio.leer_pagina(GRUPO, "CONSULTAS DIA", raiz=self.raw)
        nuevo = formato.aplicar_operaciones(doc, [
            formato.Operacion("modificar", "Reglas de negocio", "Semáforo triage II",
                              "Verde de 0 a 20 minutos, amarillo de 21 a 40 y rojo de 41 o más."),
            formato.Operacion("eliminar", "Observaciones", "Huecos en el semáforo"),
        ], resumen="Se ajustó el semáforo de triage II y se quitó una observación resuelta.", fecha="2026-10-21")
        repositorio.guardar_pagina(GRUPO, nuevo, raiz=self.raw)
        segundo = reindexar_grupo(GRUPO, cliente, project_root=self.tmp, embedder=embedder, verbose=False)
        self.assertEqual(segundo["errores"], [])
        # Una observación menos y una entrada más en el historial.
        self.assertEqual(segundo["chunks"], primero["chunks"])
        # Incremental: solo se embeben los chunks que cambiaron (regla, historial,
        # resumen de la página y la observación que se corrió de posición).
        self.assertEqual(primero["embebidos"], primero["chunks"])
        self.assertTrue(0 < segundo["embebidos"] <= 6, segundo["embebidos"])

        despues = self._por_grupo(cliente)
        self.assertEqual(despues[OTRO], antes[OTRO])  # el otro grupo no se toca
        textos = " ".join(t for t in despues[GRUPO].values() if t.startswith("Tablero: CONSULTAS DIA "))
        self.assertIn("Verde de 0 a 20 minutos", textos)
        self.assertNotIn("Verde de 0 a 15 minutos, amarillo de 16 a 30 y rojo de 31 o más. ", textos + " ")
        self.assertNotIn("Huecos en el semáforo", textos)
        self.assertIn("2026-10-21 · v2", textos)

        # Archivos derivados coherentes con Qdrant y con ambos grupos.
        chunks = json.loads(self.rutas["chunks"].read_text(encoding="utf-8"))
        self.assertEqual({c["id"] for c in chunks if c["metadata"]["source_group"] == GRUPO}, set(despues[GRUPO]))
        self.assertEqual({c["id"] for c in chunks if c["metadata"]["source_group"] == OTRO}, set(despues[OTRO]))
        catalogo = json.loads(self.rutas["catalogo"].read_text(encoding="utf-8"))
        self.assertEqual({t["source_group"] for t in catalogo["dashboards"]}, {GRUPO, OTRO})
        self.assertTrue((self.rutas["semantic_documents"] / OTRO / "1 Documentacion Otro.json").exists())

        # Sin cambios: nada que embeber y Qdrant queda igual.
        tercero = reindexar_grupo(GRUPO, cliente, project_root=self.tmp, embedder=embedder, verbose=False)
        self.assertEqual((tercero["errores"], tercero["embebidos"]), ([], 0))
        self.assertEqual(self._por_grupo(cliente), despues)

        # Reindexación completa (no incremental): mismo resultado.
        completo = reindexar_grupo(GRUPO, cliente, project_root=self.tmp, embedder=embedder,
                                   verbose=False, incremental=False)
        self.assertEqual(completo["embebidos"], completo["chunks"])
        self.assertEqual(self._por_grupo(cliente), despues)

    def test_error_de_formato_no_toca_nada(self):
        cliente = self._cliente()
        embedder = EmbedderFalso()
        reindexar_grupo(GRUPO, cliente, project_root=self.tmp, embedder=embedder, verbose=False)
        antes = self._por_grupo(cliente)
        chunks_antes = self.rutas["chunks"].read_text(encoding="utf-8")

        (self.raw / GRUPO / "md" / "rota.md").write_text("Texto fuera de sección.\n", encoding="utf-8")
        resultado = reindexar_grupo(GRUPO, cliente, project_root=self.tmp, embedder=embedder, verbose=False)
        self.assertTrue(resultado["errores"])
        self.assertIn("rota.md", " ".join(resultado["errores"]))
        self.assertEqual(self._por_grupo(cliente), antes)
        self.assertEqual(self.rutas["chunks"].read_text(encoding="utf-8"), chunks_antes)

    def test_grupo_inexistente_y_cliente_propio(self):
        resultado = reindexar_grupo("no_existe", project_root=self.tmp, embedder=EmbedderFalso(), verbose=False)
        self.assertTrue(resultado["errores"])
        # Sin cliente: abre y cierra el suyo (la carpeta queda libre).
        resultado = reindexar_grupo(GRUPO, project_root=self.tmp, embedder=EmbedderFalso(), verbose=False)
        self.assertEqual(resultado["errores"], [])
        cliente = self._cliente()
        self.assertEqual(len(self._por_grupo(cliente)[GRUPO]), resultado["chunks"])


# ============================================================
# Raíz de los .md configurable (p. ej. docs_tableros/<grupo>/md)
# ============================================================

class TestRaizMarkdown(_ConProyecto):

    def _mover_md(self, con_manifest=False):
        """Saca las páginas .md de data/raw a docs_tableros/<grupo>/md."""
        raiz_md = self.tmp / "docs_tableros"
        (raiz_md / GRUPO).mkdir(parents=True)
        shutil.move(str(self.raw / GRUPO / "md"), str(raiz_md / GRUPO / "md"))
        if con_manifest:
            manifest = json.loads((self.raw / GRUPO / "manifest.json").read_text(encoding="utf-8"))
            manifest["workspace"] = "Workspace de los md"
            (raiz_md / GRUPO / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        return raiz_md

    def _extraer(self, **kwargs):
        from src.ingestion.document_parser import process_all_documents
        from src.semantic.document_normalizer import normalize_all_documents

        with _silencio():
            extraccion = process_all_documents(self.raw, self.rutas["processed"], **kwargs)
            normalize_all_documents(self.rutas["processed"], self.rutas["semantic_documents"])
        documentos = [json.loads(p.read_text(encoding="utf-8"))
                      for p in sorted((self.rutas["semantic_documents"] / GRUPO).rglob("*.json"))]
        return extraccion, documentos

    def test_md_fuera_de_data_raw_usa_el_manifest_de_data_raw(self):
        raiz_md = self._mover_md()
        extraccion, documentos = self._extraer(md_root=raiz_md)
        self.assertEqual(extraccion["errors"], [])
        self.assertEqual([Path(p).name for p in extraccion["ignored_docx"]], [DOCX_VIEJO])
        self.assertEqual(len(documentos), len(list((raiz_md / GRUPO / "md").glob("*.md"))))
        consultas = next(d for d in documentos if d["source_file"] == "consultas_dia.md")
        self.assertEqual(consultas["semantic_model"], MODELO)
        self.assertEqual(consultas["workspace"], "Gestion Clinica")
        self.assertEqual(Path(consultas["manifest_path"]), Path(GRUPO) / "manifest.json")
        self.assertEqual(Path(consultas["relative_path"]), Path(GRUPO) / "md" / "consultas_dia.md")
        general = next(d for d in documentos if d["source_file"] == "_general.md")
        self.assertEqual(general["dashboards"][0]["name"], MODELO)

    def test_manifest_junto_a_los_md_y_docs_md_root(self):
        raiz_md = self._mover_md(con_manifest=True)
        anterior = os.environ.get("DOCS_MD_ROOT")
        os.environ["DOCS_MD_ROOT"] = str(raiz_md)
        try:
            extraccion, documentos = self._extraer()
        finally:
            if anterior is None:
                os.environ.pop("DOCS_MD_ROOT", None)
            else:
                os.environ["DOCS_MD_ROOT"] = anterior
        self.assertEqual(extraccion["errors"], [])
        self.assertTrue(documentos)
        self.assertEqual({d["workspace"] for d in documentos}, {"Workspace de los md"})

    def test_reindexar_con_md_root(self):
        raiz_md = self._mover_md()
        resultado = reindexar_grupo(GRUPO, project_root=self.tmp, md_root=raiz_md,
                                    embedder=EmbedderFalso(), verbose=False)
        self.assertEqual(resultado["errores"], [])
        self.assertEqual(resultado["documentos"], len(list((raiz_md / GRUPO / "md").glob("*.md"))))
        self.assertEqual([Path(p).name for p in resultado["docx_ignorados"]], [DOCX_VIEJO])
        self.assertTrue((self.rutas["semantic_documents"] / GRUPO / "md" / "consultas_dia.json").exists())


# ============================================================
# 5. Recuperación con el retriever del proyecto (modelo real)
# ============================================================

class TestRecuperacion(_ConProyecto):
    PREGUNTA = "¿cuál es el semáforo de triage II en consultas?"

    @classmethod
    def setUpClass(cls):
        try:
            from sentence_transformers import SentenceTransformer

            from src.rag.embedding_indexer import MODEL_NAME

            cls.modelo = SentenceTransformer(MODEL_NAME)
        except Exception as exc:  # modelo no disponible sin red
            raise unittest.SkipTest(f"Modelo de embeddings no disponible localmente: {exc}")

    def _top(self, retriever, pregunta, limite=5):
        with _silencio():
            return retriever.search_general(pregunta, limit=limite)

    def _regla_semaforo(self, retriever):
        """La regla del semáforo llega entre las 3 primeras fuentes (el
        sintetizador recibe 5) y es la primera regla de negocio. El primer
        puesto suele ser el visual «Leyendas de oportunidad», que solo
        menciona los rangos."""
        resultados = self._top(retriever, self.PREGUNTA)
        reglas = [r for r in resultados[:3] if r["chunk_type"] == "business_rule"]
        self.assertTrue(reglas, [(r["chunk_type"], r["text"][:80]) for r in resultados])
        self.assertEqual(reglas[0]["dashboard"], "CONSULTAS DIA")
        self.assertEqual(reglas[0]["source_file"], "consultas_dia.md")
        return reglas[0]

    def test_respuesta_refleja_el_texto_nuevo_tras_reindexar(self):
        from src.rag.retriever import HybridRetriever, active_retriever

        for grupo in (OTRO, GRUPO):
            resultado = reindexar_grupo(grupo, project_root=self.tmp, embedder=self.modelo, verbose=False)
            self.assertEqual(resultado["errores"], [], grupo)

        retriever = HybridRetriever(self.rutas["qdrant"])
        self.addCleanup(retriever.close)
        self.assertIs(active_retriever(self.rutas["qdrant"]), retriever)

        regla = self._regla_semaforo(retriever)
        self.assertIn("Semáforo triage II: Verde de 0 a 15 minutos", regla["text"])

        # Pregunta por otra página del mismo tablero.
        resultados = self._top(retriever, "¿cuál es el semáforo de interconsultas del día?")
        self.assertIn("Semáforo de interconsultas", " ".join(r["text"] for r in resultados[:2]))

        # Cambio aprobado en el .md y reindexación usando el cliente de la app
        # (se detecta solo: la carpeta ya está abierta por el retriever).
        doc = repositorio.leer_pagina(GRUPO, "CONSULTAS DIA", raiz=self.raw)
        nuevo = formato.aplicar_operaciones(doc, [formato.Operacion(
            "modificar", "Reglas de negocio", "Semáforo triage II",
            "Verde de 0 a 20 minutos, amarillo de 21 a 40 y rojo de 41 o más.")],
            resumen="Se ajustó el umbral verde del semáforo de triage II.", fecha="2026-10-21")
        repositorio.guardar_pagina(GRUPO, nuevo, raiz=self.raw)
        resultado = reindexar_grupo(GRUPO, project_root=self.tmp, verbose=False)
        self.assertEqual(resultado["errores"], [])
        self.assertTrue(resultado["retriever_actualizado"])

        regla = self._regla_semaforo(retriever)
        self.assertIn("Semáforo triage II: Verde de 0 a 20 minutos", regla["text"])
        for item in self._top(retriever, self.PREGUNTA):
            self.assertNotIn("0 a 15 minutos, amarillo de 16 a 30", item["text"]
                             if item["dashboard"] == "CONSULTAS DIA" else "")
        # El otro grupo sigue en el índice.
        self.assertIn(OTRO, {item["source_group"] for item in retriever._corpus})


if __name__ == "__main__":
    unittest.main()
